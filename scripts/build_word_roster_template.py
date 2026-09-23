#!/usr/bin/env python3
"""Build the Arabic, machine-readable Word roster template.

The visible form is intentionally simple. Machine tags live in Word content
control metadata so administrators never need to type or understand them.
"""

from __future__ import annotations

import copy
import io
import re
import sys
import zipfile
import zlib
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Inches, Pt, RGBColor
from lxml import etree


W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
NS = {"w": W, "wp": WP}

NAVY = "12395F"
NAVY_DARK = "0B2A4A"
PALE_BLUE = "EAF2F8"
PALE_GOLD = "FFF8E6"
LIGHT_BORDER = "D9E2EC"
TEXT = "12263A"
MUTED = "5E7184"


def set_run_font(run, size=10.5, bold=False, color=TEXT):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Arial")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rtl = OxmlElement("w:rtl")
    rtl.set(qn("w:val"), "1")
    run._element.get_or_add_rPr().append(rtl)


def set_paragraph_rtl(paragraph, alignment=WD_ALIGN_PARAGRAPH.RIGHT, before=0, after=0):
    paragraph.alignment = alignment
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.0
    ppr = paragraph._p.get_or_add_pPr()
    bidi = ppr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        ppr.append(bidi)
    bidi.set(qn("w:val"), "1")


def set_cell_shading(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tcpr = cell._tc.get_or_add_tcPr()
    mar = tcpr.find(qn("w:tcMar"))
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tcpr.append(mar)
    for name, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            mar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_BORDER, size="6"):
    tblpr = table._tbl.tblPr
    borders = tblpr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def make_table_rtl(table):
    tblpr = table._tbl.tblPr
    bidi = tblpr.find(qn("w:bidiVisual"))
    if bidi is None:
        bidi = OxmlElement("w:bidiVisual")
        tblpr.insert(0, bidi)
    bidi.set(qn("w:val"), "1")


def prevent_row_split(row):
    trpr = row._tr.get_or_add_trPr()
    if trpr.find(qn("w:cantSplit")) is None:
        trpr.append(OxmlElement("w:cantSplit"))


def clear_cell(cell):
    cell.text = ""
    p = cell.paragraphs[0]
    set_paragraph_rtl(p)
    return p


def label_cell(cell, text):
    p = clear_cell(cell)
    set_cell_shading(cell, NAVY)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    r = p.add_run(text)
    set_run_font(r, size=9.5, bold=True, color="FFFFFF")


def field_cell(cell, tag):
    p = clear_cell(cell)
    # Keep the form/table direction Arabic (RTL), but center every value that
    # the administrator types or selects. This avoids values drifting to the
    # left or right depending on whether they contain Arabic text or digits.
    set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_shading(cell, "FFFFFF")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    r = p.add_run("{{" + tag + "}}")
    set_run_font(r, size=10.5, color=TEXT)


def title_row(table, text):
    merged = table.cell(0, 0).merge(table.cell(0, len(table.columns) - 1))
    p = clear_cell(merged)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_cell_shading(merged, NAVY_DARK)
    r = p.add_run(text)
    set_run_font(r, size=11, bold=True, color="FFFFFF")


def extract_placeholders(source):
    with zipfile.ZipFile(source) as zf:
        return zf.read("word/media/image1.png"), zf.read("word/media/image2.png")


def add_picture_box(cell, label, image_bytes, tag, private=False):
    clear_cell(cell)
    set_cell_shading(cell, PALE_GOLD if private else PALE_BLUE)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    label_p = cell.paragraphs[0]
    label_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label_r = label_p.add_run(label)
    set_run_font(label_r, size=9.5, bold=True, color=NAVY_DARK)

    pic_p = cell.add_paragraph()
    set_paragraph_rtl(pic_p, WD_ALIGN_PARAGRAPH.CENTER)
    pic_r = pic_p.add_run()
    inline = pic_r.add_picture(io.BytesIO(image_bytes), width=Inches(1.55))._inline
    inline.docPr.set("name", tag)
    inline.docPr.set("descr", tag)

    hint_p = cell.add_paragraph()
    set_paragraph_rtl(hint_p, WD_ALIGN_PARAGRAPH.CENTER)
    hint_r = hint_p.add_run("اضغط على رمز الصورة لاختيارها في Word على الكمبيوتر")
    set_run_font(hint_r, size=7.5, color=MUTED)


def configure_document_defaults(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.15)
    section.bottom_margin = Cm(1.15)
    section.left_margin = Cm(1.35)
    section.right_margin = Cm(1.35)
    section.header_distance = Cm(0.4)
    section.footer_distance = Cm(0.4)
    sectpr = section._sectPr
    if sectpr.find(qn("w:rtlGutter")) is None:
        sectpr.append(OxmlElement("w:rtlGutter"))

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    rpr = normal.element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn("w:ascii"), "Arial")
    rpr.get_or_add_rFonts().set(qn("w:hAnsi"), "Arial")
    rpr.get_or_add_rFonts().set(qn("w:cs"), "Arial")

    title = doc.styles["Title"]
    title.font.name = "Arial"
    title.font.color.rgb = RGBColor(0, 0, 0)
    title_ppr = title.element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)


def add_intro(doc):
    p = doc.add_paragraph()
    set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.CENTER, after=2)
    r = p.add_run("نموذج بيانات فريق البطولة")
    set_run_font(r, size=18, bold=True, color="000000")
    p.style = doc.styles["Title"]

    p = doc.add_paragraph()
    set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.CENTER, after=6)
    r = p.add_run("قائمة اللاعبين والصور")
    set_run_font(r, size=11, bold=True, color=NAVY)

    # Keep the instructions inside the same visual width as the form instead
    # of allowing one long line to stretch across the page. Two clear rows are
    # easier for non-technical administrators to scan.
    instructions = doc.add_table(rows=2, cols=1)
    instructions.alignment = WD_TABLE_ALIGNMENT.CENTER
    instructions.autofit = False
    instructions.columns[0].width = Cm(17.6)
    make_table_rtl(instructions)
    set_table_borders(instructions, color=LIGHT_BORDER, size="5")
    instruction_rows = (
        "١. اضغط داخل الخانة واكتب البيانات. لإضافة الصورة الشخصية، اضغط رمز الصورة واخترها من جهازك.",
        "٢. الصورة الشخصية تُحفظ بعد المراجعة. صورة الهوية أو الجواز تبقى داخل Word وملف PDF الخاص، ولا تُحفظ في قاعدة البيانات.",
    )
    for index, text in enumerate(instruction_rows):
        cell = instructions.cell(index, 0)
        cell.width = Cm(17.6)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, top=85, start=130, bottom=85, end=130)
        set_cell_shading(cell, "F7FAFC" if index == 0 else PALE_GOLD)
        p = clear_cell(cell)
        set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
        r = p.add_run(text)
        set_run_font(r, size=10.5, bold=True if index == 0 else False, color=TEXT)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(3)


def add_team_table(doc):
    table = doc.add_table(rows=4, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    make_table_rtl(table)
    set_table_borders(table)
    widths = [Cm(2.6), Cm(6.2), Cm(2.6), Cm(6.2)]
    for row in table.rows:
        prevent_row_split(row)
        for i, cell in enumerate(row.cells):
            cell.width = widths[i]
            set_cell_margins(cell)

    title_row(table, "بيانات الفريق")
    label_cell(table.cell(1, 0), "اسم الفريق")
    team_field = table.cell(1, 1).merge(table.cell(1, 3))
    field_cell(team_field, "TEAM_NAME")

    label_cell(table.cell(2, 0), "مسؤول الفريق")
    field_cell(table.cell(2, 1), "TEAM_ADMIN_NAME")
    label_cell(table.cell(2, 2), "الجنسية")
    field_cell(table.cell(2, 3), "TEAM_ADMIN_NATION")

    label_cell(table.cell(3, 0), "المدرب")
    field_cell(table.cell(3, 1), "COACH_NAME")
    label_cell(table.cell(3, 2), "الجنسية")
    field_cell(table.cell(3, 3), "COACH_NATION")


def add_player_table(doc, number, public_image, private_image):
    pnum = f"{number:02d}"
    table = doc.add_table(rows=5, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    make_table_rtl(table)
    set_table_borders(table)
    widths = [Cm(2.6), Cm(6.2), Cm(2.6), Cm(6.2)]
    for row in table.rows:
        prevent_row_split(row)
        for i, cell in enumerate(row.cells):
            cell.width = widths[i]
            set_cell_margins(cell, top=65, bottom=65)

    title_row(table, f"بيانات اللاعب {number}")

    label_cell(table.cell(1, 0), "الاسم الأول")
    field_cell(table.cell(1, 1), f"P{pnum}_FIRST_NAME")
    label_cell(table.cell(1, 2), "اسم العائلة")
    field_cell(table.cell(1, 3), f"P{pnum}_LAST_NAME")

    label_cell(table.cell(2, 0), "رقم القميص")
    field_cell(table.cell(2, 1), f"P{pnum}_NUMBER")
    label_cell(table.cell(2, 2), "الجنسية")
    field_cell(table.cell(2, 3), f"P{pnum}_NATION")

    label_cell(table.cell(3, 0), "حارس مرمى؟")
    goalkeeper_field = table.cell(3, 1).merge(table.cell(3, 3))
    field_cell(goalkeeper_field, f"P{pnum}_GK")

    public_cell = table.cell(4, 0).merge(table.cell(4, 1))
    private_cell = table.cell(4, 2).merge(table.cell(4, 3))
    add_picture_box(
        public_cell,
        "الصورة الشخصية العامة",
        public_image,
        f"PLAYER_{pnum}_PUBLIC",
    )
    add_picture_box(
        private_cell,
        "صورة الهوية أو الجواز الخاصة",
        private_image,
        f"PLAYER_{pnum}_PRIVATE",
        private=True,
    )
    return table


def content_prompt(tag):
    if tag == "TEAM_NAME":
        return "اكتب اسم الفريق"
    if tag in {"TEAM_ADMIN_NAME", "COACH_NAME"}:
        return "اكتب الاسم الكامل"
    if tag in {"TEAM_ADMIN_NATION", "COACH_NATION"}:
        return "اكتب الجنسية"
    if tag.endswith("_FIRST_NAME"):
        return "اكتب الاسم الأول"
    if tag.endswith("_LAST_NAME"):
        return "اكتب اسم العائلة"
    if tag.endswith("_NUMBER"):
        return "اكتب رقمًا من 1 إلى 99 مثل 3 أو 11"
    if tag.endswith("_NATION"):
        return "اكتب الجنسية"
    if tag.endswith("_GK"):
        return "اختر نعم للحارس فقط، وإلا اتركها فارغة ▼"
    return "اكتب هنا"


def arabic_alias(tag):
    if tag == "TEAM_NAME":
        return "اسم الفريق"
    if tag == "TEAM_ADMIN_NAME":
        return "اسم مسؤول الفريق"
    if tag == "TEAM_ADMIN_NATION":
        return "جنسية مسؤول الفريق"
    if tag == "COACH_NAME":
        return "اسم المدرب"
    if tag == "COACH_NATION":
        return "جنسية المدرب"
    m = re.match(r"P(\d{2})_(.+)", tag)
    if m:
        labels = {
            "FIRST_NAME": "الاسم الأول",
            "LAST_NAME": "اسم العائلة",
            "NUMBER": "رقم القميص",
            "NATION": "الجنسية",
            "GK": "حارس مرمى",
        }
        return f"اللاعب {m.group(1)} - {labels.get(m.group(2), m.group(2))}"
    if tag.startswith("PLAYER_"):
        return "صورة اللاعب" if tag.endswith("_PUBLIC") else "صورة الهوية الخاصة"
    return tag


def elem(tag, **attrs):
    node = etree.Element(f"{{{W}}}{tag}")
    for key, value in attrs.items():
        node.set(f"{{{W}}}{key}", str(value))
    return node


def add_common_sdt_props(sdtpr, tag, picture=False):
    sdtpr.append(elem("alias", val=arabic_alias(tag)))
    sdtpr.append(elem("tag", val=tag))
    # Stable IDs make the generated DOCX reproducible while keeping every
    # content control distinct.
    sdtpr.append(elem("id", val=1000000 + zlib.crc32(tag.encode("utf-8"))))
    if picture:
        # A picture SDT must explicitly mark its current content as a
        # placeholder. Without showingPlcHdr Word treats the drawing like an
        # ordinary selected image and only opens the Picture Format ribbon.
        placeholder = etree.SubElement(sdtpr, f"{{{W}}}placeholder")
        placeholder.append(elem("docPart", val="DefaultPlaceholder_1082065158"))
        sdtpr.append(elem("showingPlcHdr"))
        sdtpr.append(elem("picture"))


def wrap_text_controls(root):
    for text_node in list(root.xpath(".//w:t", namespaces=NS)):
        raw = text_node.text or ""
        match = re.fullmatch(r"\{\{([A-Z0-9_]+)\}\}", raw)
        if not match:
            continue
        tag = match.group(1)
        run = text_node.getparent()
        if run.tag != f"{{{W}}}r":
            continue
        parent = run.getparent()
        sdt = etree.Element(f"{{{W}}}sdt")
        sdtpr = etree.SubElement(sdt, f"{{{W}}}sdtPr")
        add_common_sdt_props(sdtpr, tag)

        if tag.endswith("_GK"):
            dd = etree.SubElement(sdtpr, f"{{{W}}}dropDownList")
            for display, value in (
                ("اختر نعم للحارس فقط، وإلا اتركها فارغة ▼", ""),
                ("نعم — حارس مرمى", "نعم"),
            ):
                item = etree.SubElement(dd, f"{{{W}}}listItem")
                item.set(f"{{{W}}}displayText", display)
                item.set(f"{{{W}}}value", value)
        else:
            sdtpr.append(elem("text"))
            sdtpr.append(elem("showingPlcHdr"))

        content = etree.SubElement(sdt, f"{{{W}}}sdtContent")
        new_run = copy.deepcopy(run)
        new_texts = new_run.xpath(".//w:t", namespaces=NS)
        if new_texts:
            new_texts[0].text = content_prompt(tag)
        content.append(new_run)
        parent.replace(run, sdt)


def wrap_picture_controls(root):
    for inline in list(root.xpath(".//wp:inline", namespaces=NS)):
        docpr = inline.find(f"{{{WP}}}docPr")
        if docpr is None:
            continue
        tag = docpr.get("name", "")
        if not re.fullmatch(r"PLAYER_\d{2}_(PUBLIC|PRIVATE)", tag):
            continue
        run = inline
        while run is not None and run.tag != f"{{{W}}}r":
            run = run.getparent()
        if run is None:
            continue
        parent = run.getparent()
        sdt = etree.Element(f"{{{W}}}sdt")
        sdtpr = etree.SubElement(sdt, f"{{{W}}}sdtPr")
        add_common_sdt_props(sdtpr, tag, picture=True)
        content = etree.SubElement(sdt, f"{{{W}}}sdtContent")
        content.append(copy.deepcopy(run))
        parent.replace(run, sdt)


def force_input_alignment(root):
    """Force every data-entry paragraph to remain centered in desktop Word.

    Word can retain paragraph properties from an earlier placeholder when an
    inline content control is replaced. Re-applying w:jc after wrapping the
    controls makes the intended alignment explicit in the final OOXML.
    """
    for sdt in root.xpath(".//w:sdt", namespaces=NS):
        if sdt.xpath("./w:sdtPr/w:picture", namespaces=NS):
            continue
        paragraph = sdt.getparent()
        while paragraph is not None and paragraph.tag != f"{{{W}}}p":
            paragraph = paragraph.getparent()
        if paragraph is None:
            continue
        ppr = paragraph.find(f"{{{W}}}pPr")
        if ppr is None:
            ppr = etree.Element(f"{{{W}}}pPr")
            paragraph.insert(0, ppr)
        for old in list(ppr.findall(f"{{{W}}}jc")):
            ppr.remove(old)
        ppr.append(elem("jc", val="center"))
        bidi = ppr.find(f"{{{W}}}bidi")
        if bidi is None:
            bidi = elem("bidi", val="1")
            ppr.append(bidi)
        else:
            bidi.set(f"{{{W}}}val", "1")


def patch_ooxml(input_docx, output_docx):
    with zipfile.ZipFile(input_docx) as zin:
        entries = {name: zin.read(name) for name in zin.namelist()}

    root = etree.fromstring(entries["word/document.xml"])
    wrap_text_controls(root)
    wrap_picture_controls(root)
    force_input_alignment(root)
    entries["word/document.xml"] = etree.tostring(
        root, xml_declaration=True, encoding="UTF-8", standalone="yes"
    )

    settings = etree.fromstring(entries["word/settings.xml"])
    view = settings.find(f"{{{W}}}view")
    if view is None:
        view = elem("view", val="print")
        settings.insert(0, view)
    else:
        view.set(f"{{{W}}}val", "print")
    zoom = settings.find(f"{{{W}}}zoom")
    if zoom is None:
        zoom = elem("zoom", percent="95")
        settings.insert(1, zoom)
    else:
        zoom.set(f"{{{W}}}percent", "95")
    theme_lang = settings.find(f"{{{W}}}themeFontLang")
    if theme_lang is None:
        theme_lang = elem("themeFontLang")
        settings.append(theme_lang)
    theme_lang.set(f"{{{W}}}val", "ar-SA")
    theme_lang.set(f"{{{W}}}bidi", "ar-SA")
    entries["word/settings.xml"] = etree.tostring(
        settings, xml_declaration=True, encoding="UTF-8", standalone="yes"
    )

    with zipfile.ZipFile(output_docx, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in entries.items():
            zout.writestr(name, data)


def build(source, output):
    public_image, private_image = extract_placeholders(source)
    doc = Document()
    configure_document_defaults(doc)
    add_intro(doc)
    add_team_table(doc)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)
    add_player_table(doc, 1, public_image, private_image)
    doc.add_page_break()

    for number in range(2, 16):
        add_player_table(doc, number, public_image, private_image)
        if number < 15:
            if number % 2 == 1:
                doc.add_page_break()
            else:
                spacer = doc.add_paragraph()
                spacer.paragraph_format.space_after = Pt(6)

    for paragraph in doc.paragraphs:
        if paragraph._p.get_or_add_pPr().find(qn("w:bidi")) is None:
            set_paragraph_rtl(paragraph, paragraph.alignment or WD_ALIGN_PARAGRAPH.RIGHT)

    temp = output.with_suffix(".building.docx")
    doc.save(temp)
    patch_ooxml(temp, output)
    temp.unlink()


def main():
    if len(sys.argv) != 3:
        raise SystemExit("usage: build_word_roster_template.py SOURCE.docx OUTPUT.docx")
    build(Path(sys.argv[1]), Path(sys.argv[2]))


if __name__ == "__main__":
    main()
